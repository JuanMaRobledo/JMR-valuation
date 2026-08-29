"""Lee un CompanyInputs directamente del Excel 'Modelo JMR' -- sin pasar por un CSV.

Requiere que el archivo ya haya sido abierto y recalculado en Excel al menos
una vez (para que las formulas tengan un valor en cache; openpyxl no calcula
formulas, solo lee lo que Excel ya calculo). Si acabas de editar 'Input sheet'
y no viste los numeros actualizarse en el propio Excel, abrilo, dejalo
recalcular, guardalo, y recién despues cargalo aca.

Cubre 'Input sheet', 'Resumen de Valoracion' (tipo de empresa y margen de
seguridad), 'Trailing Valuation' (multiplos historicos), 'Operating lease
converter' / 'R& D converter' (si los flags correspondientes estan en 'Yes'),
y calcula los ratios historicos desde 'Income Statement' / 'Cash Flow
Statement' / 'Balance Sheet'.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl

from jmr_valuation.io.inputs import CompanyInputs
from jmr_valuation.models.financials_multiples import historical_ratios_from_actuals

# 'Resumen de Valoracion'!G3 usa nombres en espanol con tildes; models.blend.COMPANY_TYPES
# los usa sin tildes (ASCII) para no depender del encoding del CSV.
_COMPANY_TYPE_FROM_EXCEL = {
    "Crecimiento": "Crecimiento", "Madura": "Madura", "Genérico": "Generico",
    "Defensiva": "Defensiva", "Cíclica/Commodity": "Ciclica/Commodity",
    "Intensiva en Capital": "Intensiva en Capital", "Financiera": "Financiera",
    "Infraestructura": "Infraestructura", "REIT/Inmobiliaria": "REIT/Inmobiliaria",
    "Software": "Software",
}


def _yes(value) -> bool:
    return str(value).strip().lower() in ("yes", "si", "sí", "true")


def _num(value, default: float = 0.0) -> float:
    return float(value) if isinstance(value, (int, float)) else default


def load_company_inputs_from_excel(xlsx_path: str | Path) -> CompanyInputs:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ins = wb["Input sheet"]

    capitalize_rd = _yes(ins["B17"].value)
    has_operating_leases = _yes(ins["B18"].value)
    has_employee_options = _yes(ins["B38"].value)

    kwargs = dict(
        ticker=str(ins["B1"].value or ""), company_name=str(ins["B5"].value or ins["B1"].value or ""),
        country_of_incorporation=str(ins["B8"].value or ""), industry_us=str(ins["B9"].value or ""),
        industry_global=str(ins["B10"].value or ""),
        revenue_ltm=_num(ins["B12"].value), revenue_prior_10k=_num(ins["C12"].value),
        years_since_last_10k=_num(ins["D12"].value, 0.5),
        ebit_ltm=_num(ins["B13"].value), ebit_prior_10k=_num(ins["C13"].value),
        interest_expense_ltm=_num(ins["B14"].value), interest_expense_prior_10k=_num(ins["C14"].value),
        book_value_equity_ltm=_num(ins["B15"].value), book_value_equity_prior_10k=_num(ins["C15"].value),
        book_value_debt_ltm=_num(ins["B16"].value), book_value_debt_prior_10k=_num(ins["C16"].value),
        cash_ltm=_num(ins["B19"].value), cash_prior_10k=_num(ins["C19"].value),
        cross_holdings_ltm=_num(ins["B20"].value), cross_holdings_prior_10k=_num(ins["C20"].value),
        minority_interests=_num(ins["B21"].value), shares_outstanding=_num(ins["B22"].value),
        current_price=_num(ins["B23"].value), effective_tax_rate=_num(ins["B24"].value),
        marginal_tax_rate=_num(ins["B25"].value, 0.25),
        capitalize_rd=capitalize_rd, has_operating_leases=has_operating_leases,
        has_employee_options=has_employee_options,
        revenue_growth_next_year=_num(ins["B27"].value), revenue_growth_years_2_to_5=_num(ins["B29"].value),
        target_ebit_margin=_num(ins["B30"].value), year_of_margin_convergence=int(_num(ins["B31"].value, 5)),
        sales_to_capital_years_1_5=_num(ins["B32"].value, 2.0),
        sales_to_capital_years_6_10=_num(ins["B33"].value, 2.0),
        riskfree_rate=_num(ins["B35"].value), initial_cost_of_capital=_num(ins["B36"].value),
        n_options_outstanding=_num(ins["B39"].value), avg_strike_price=_num(ins["B40"].value),
        avg_option_maturity=_num(ins["B41"].value), stdev_stock_price=_num(ins["B42"].value),
        nol_carryforward=_num(ins["B63"].value),
    )

    if "Resumen de Valoración" in wb.sheetnames:
        resumen = wb["Resumen de Valoración"]
        excel_type = str(resumen["G3"].value or "").strip()
        kwargs["company_type"] = _COMPANY_TYPE_FROM_EXCEL.get(excel_type, "Generico")
        kwargs["margin_of_safety"] = _num(resumen["G4"].value, 0.35)

    if "Trailing Valuation" in wb.sheetnames:
        tv = wb["Trailing Valuation"]
        # Columnas I,J,K = los ultimos 3 anios completos reportados; K es el mas reciente (anio -1).
        multiple_rows = {"ev_fcff": 24, "p_ocf": 15, "pe": 13, "p_fcfe": 16, "ev_ebitda": 21}
        for prefix, row in multiple_rows.items():
            kwargs[f"hist_multiple_{prefix}_y1"] = _num(tv.cell(row=row, column=11).value)  # K
            kwargs[f"hist_multiple_{prefix}_y2"] = _num(tv.cell(row=row, column=10).value)  # J
            kwargs[f"hist_multiple_{prefix}_y3"] = _num(tv.cell(row=row, column=9).value)   # I

    # Si la celda J19 (override del bloque Base) tiene un valor en la hoja del
    # multiplo, ese es el multiplo a usar -- igual que
    # 'EVFCFF'!F19=IF(J19<>"",J19,MEDIAN(...)) en el Excel.
    multiple_sheet_prefix = {"EVFCFF": "ev_fcff", "POCF": "p_ocf", "PE": "pe",
                              "PFCFE": "p_fcfe", "EVEBITDA": "ev_ebitda"}
    for sheet_name, prefix in multiple_sheet_prefix.items():
        if sheet_name not in wb.sheetnames:
            continue
        j19 = wb[sheet_name]["J19"].value
        if j19 is not None and j19 != "":
            kwargs[f"override_{prefix}"] = True
            kwargs[f"multiple_{prefix}"] = _num(j19)

    if "Income Statement" in wb.sheetnames:
        inc = wb["Income Statement"]
        # 'Valuation output'!C49:C52 -- LTM / promedio 3a / promedio 5a / promedio 10a
        # del margen operativo *sin ajustar* (fila 13 de Income Statement).
        margins_3y = [v for v in (inc.cell(row=13, column=c).value for c in range(9, 12)) if isinstance(v, (int, float))]
        margins_5y = [v for v in (inc.cell(row=13, column=c).value for c in range(7, 12)) if isinstance(v, (int, float))]
        margins_10y = [v for v in (inc.cell(row=13, column=c).value for c in range(2, 12)) if isinstance(v, (int, float))]

        kwargs["ebit_margin_ltm"] = _num(inc["L13"].value)
        kwargs["ebit_margin_avg_3y"] = sum(margins_3y) / len(margins_3y) if margins_3y else 0.0
        kwargs["ebit_margin_avg_5y"] = sum(margins_5y) / len(margins_5y) if margins_5y else 0.0
        kwargs["ebit_margin_avg_10y"] = sum(margins_10y) / len(margins_10y) if margins_10y else 0.0

    if has_operating_leases and "Operating lease converter" in wb.sheetnames:
        olc = wb["Operating lease converter"]
        kwargs["lease_expense_current_year"] = _num(olc["E5"].value)
        kwargs["lease_commitment_y1"] = _num(olc["B8"].value)
        kwargs["lease_commitment_y2"] = _num(olc["B9"].value)
        kwargs["lease_commitment_y3"] = _num(olc["B10"].value)
        kwargs["lease_commitment_y4"] = _num(olc["B11"].value)
        kwargs["lease_commitment_y5"] = _num(olc["B12"].value)
        kwargs["lease_commitment_y6_plus"] = _num(olc["B13"].value)

    if capitalize_rd and "R& D converter" in wb.sheetnames:
        rdc = wb["R& D converter"]
        kwargs["rd_amortization_years"] = int(_num(rdc["F7"].value, 3))
        kwargs["rd_expense_current_year"] = _num(rdc["F8"].value)
        # B12..B20 = anio -1 hasta -9 (fila 12 = anio -1, ver hoja 'R& D converter')
        for i, row in enumerate(range(12, 21), start=1):
            kwargs[f"rd_expense_year_minus_{i}"] = _num(rdc.cell(row=row, column=2).value)

    if all(s in wb.sheetnames for s in ("Income Statement", "Cash Flow Statement", "Balance Sheet")):
        _fill_historical_ratios(wb, kwargs)

    if "Dividendos" in wb.sheetnames:
        div = wb["Dividendos"]
        kwargs["dividend_per_share_ltm"] = _num(div["K4"].value)

    return CompanyInputs(**kwargs)


def _fill_historical_ratios(wb, kwargs: dict) -> None:
    inc = wb["Income Statement"]
    cf = wb["Cash Flow Statement"]
    bs = wb["Balance Sheet"]

    def row_hijk(ws, row):
        return [_num(ws.cell(row=row, column=c).value) for c in (8, 9, 10, 11)]  # H,I,J,K

    revenue_h, *revenue_ijk = row_hijk(inc, 3)
    _, *ebit_ijk = row_hijk(inc, 12)
    _, *interest_ijk = row_hijk(inc, 18)
    _, *da_ijk = row_hijk(cf, 4)
    capex_hijk = row_hijk(cf, 15)
    capex_ijk = [abs(v) for v in capex_hijk[1:]]  # magnitud positiva, ver docstring de financials_multiples
    _, *nb1_ijk = [a + b for a, b in zip(row_hijk(cf, 25), row_hijk(cf, 28))]
    shares_h, *shares_ijk = row_hijk(inc, 26)

    bs5, bs10, bs20, bs21, bs24 = (row_hijk(bs, r) for r in (5, 10, 20, 21, 24))
    nwc_term = [(bs10[i] - bs5[i]) - (bs24[i] - bs20[i] - bs21[i]) for i in range(4)]
    nwc_change_ijk = [nwc_term[i] - nwc_term[i - 1] for i in (1, 2, 3)]

    try:
        ratios = historical_ratios_from_actuals(
            revenues=revenue_ijk, prior_year_revenue=revenue_h,
            ebit=ebit_ijk, interest_other=interest_ijk, da=da_ijk, capex=capex_ijk,
            nwc_change=nwc_change_ijk, net_borrowing=nb1_ijk,
            shares_diluted=shares_ijk, prior_year_shares_diluted=shares_h,
        )
    except (ZeroDivisionError, ValueError):
        return

    kwargs["hist_interest_pct_of_ebit"] = ratios.interest_pct_of_ebit
    kwargs["hist_da_pct_of_revenue"] = ratios.da_pct_of_revenue
    kwargs["hist_capex_pct_of_revenue"] = ratios.capex_pct_of_revenue
    kwargs["hist_nwc_pct_of_revenue_growth"] = ratios.nwc_change_pct_of_revenue_growth
    kwargs["hist_net_borrowing_pct_of_revenue"] = ratios.net_borrowing_pct_of_revenue
    kwargs["hist_shares_growth_rate"] = ratios.shares_growth_rate
