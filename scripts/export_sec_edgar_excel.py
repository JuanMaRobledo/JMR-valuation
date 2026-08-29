#!/usr/bin/env python
"""Genera un Excel de revision con los estados financieros crudos de SEC EDGAR
y los ratios/margenes ya calculados por FORMULA de Excel (no numeros fijos) --
para poder auditar/corregir un numero puntual y ver el ratio recalcularse solo,
antes de subir los datos al dashboard.

Uso:
    python scripts/export_sec_edgar_excel.py INTU
    python scripts/export_sec_edgar_excel.py INTU --price 650 --riskfree 0.045 --wacc 0.09

Genera dos archivos en --out-dir (por defecto data/):
    <TICKER>_sec_edgar.xlsx  -- para revisar (2 hojas, con formulas)
    <TICKER>_sec_edgar.csv   -- listo para subir tal cual con el boton
                                "...o subir un CSV" del dashboard (mismos
                                numeros que las celdas de la hoja 'Inputs
                                dashboard' del xlsx, pero ya resueltos).

IMPORTANTE: el .xlsx que genera este script NO es el 'Modelo JMR' original
(ese tiene sus propias hojas -- 'Input sheet', 'Trailing Valuation', etc. --
que este script no reproduce). No lo subas con el boton "...o subir el Excel
del modelo (.xlsx)"; usa el .csv que se genera al lado, o copia los valores
de la hoja 'Inputs dashboard' a mano.

Lo que SEC EDGAR no tiene y hay que completar a mano (parametros --price,
--riskfree, --wacc si no se pasan quedan en 0 y marcados como pendientes en
el xlsx): precio de mercado, tasa libre de riesgo, costo de capital (WACC),
y los multiplos historicos (P/E, EV/EBITDA, etc. -- dependen del precio en
cada fecha, que EDGAR no tiene).
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from jmr_valuation.io.inputs import save_company_inputs
from jmr_valuation.io.sec_edgar_client import SecEdgarClient, SecEdgarError
from jmr_valuation.io.sec_edgar_loader import (
    _annual_instant_rows,
    _annual_rows,
    _concept_rows,
    _instant_as_of,
    _latest_instant,
    _quarters_are_contiguous,
    _quarterly_rows,
    _resolve_country,
    load_company_inputs_from_sec_edgar,
)


def _last_four_contiguous_quarters(rows: list[dict] | None) -> list[dict]:
    """Como _ltm_value: si faltan trimestres o los ultimos 4 disponibles no
    son consecutivos (comun quando el ultimo trimestre fiscal nunca se taggea
    solo, p.ej. Intuit), no hay un LTM real que mostrar -- se devuelve vacio y
    el llamador cae al ultimo 10-K, igual que hace el loader."""
    quarters = _quarterly_rows(rows)
    if len(quarters) >= 4 and _quarters_are_contiguous(quarters[-4:]):
        return quarters[-4:]
    return []

FILL_INPUT = PatternFill("solid", fgColor="FFF3CD")   # amarillo -- hay que completarlo a mano
FILL_FORMULA = PatternFill("solid", fgColor="EAF2FB")  # celeste -- se calcula solo
FILL_HEADER = PatternFill("solid", fgColor="1B2A4A")
BOLD = Font(bold=True)
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _year_values(annual_rows: list[dict], ends: list[str]) -> dict[str, float]:
    by_end = {r["end"]: r["val"] / 1_000_000 for r in annual_rows}
    return {e: by_end.get(e) for e in ends}


def _cagr_formula(sheet: str, row: int, last_col: str, n_years: int, first_available_col: str) -> str:
    """CAGR de 3 anios (ultimo valor vs. el de N anios atras) -- consistente
    con jmr_valuation.io.sec_edgar_loader._cagr(), que usa un CAGR de 3 anios
    (no un promedio de tasas anuales, que diluiria una tendencia reciente con
    datos de hace una decada) como estimador ingenuo de partida."""
    back_col_idx = max(_col_index(first_available_col), _col_index(last_col) - n_years)
    back_col = get_column_letter(back_col_idx)
    years = _col_index(last_col) - back_col_idx
    if years < 1:
        return "0"
    return f"=IF(OR('{sheet}'!{back_col}{row}=\"\",'{sheet}'!{back_col}{row}<=0),0,('{sheet}'!{last_col}{row}/'{sheet}'!{back_col}{row})^(1/{years})-1)"


def _col_index(col_letter: str) -> int:
    from openpyxl.utils import column_index_from_string
    return column_index_from_string(col_letter)


def _write_row(ws, row: int, label: str, ends: list[str], values: dict[str, float | None],
                start_col: int = 3, number_format: str = "#,##0.0") -> None:
    ws.cell(row=row, column=2, value=label)
    for i, end in enumerate(ends):
        cell = ws.cell(row=row, column=start_col + i, value=values.get(end))
        cell.number_format = number_format


def build_workbook(ticker: str, current_price: float, riskfree_rate: float,
                    initial_cost_of_capital: float) -> tuple[openpyxl.Workbook, object]:
    client = SecEdgarClient()
    facts = client.company_facts(ticker)
    submissions = client.company_submissions(ticker)
    gaap = facts.get("facts", {}).get("us-gaap", {})

    def rows(key: str, units=("USD",)):
        return _concept_rows(gaap, key, units)

    revenue_annual = _annual_rows(rows("revenue"))
    if not revenue_annual:
        raise SecEdgarError(f"No se encontraron ingresos anuales (10-K) para {ticker} en SEC EDGAR.")
    ebit_annual = _annual_rows(rows("ebit"))

    long_ends = [r["end"] for r in revenue_annual[-10:]]
    short_ends = long_ends[-5:]
    last_10k_end = long_ends[-1]

    revenue_by_end = {r["end"]: r["val"] / 1_000_000 for r in revenue_annual}
    ebit_by_end = {r["end"]: r["val"] / 1_000_000 for r in ebit_annual}

    interest_annual = _annual_rows(rows("interest_expense"))
    da_annual = _annual_rows(rows("da"))
    capex_annual = _annual_rows(rows("capex"))
    rd_annual = _annual_rows(rows("rd"))
    tax_annual = _annual_rows(rows("tax_expense"))
    pretax_annual = _annual_rows(rows("pretax_income"))
    dividend_annual = _annual_rows(rows("dividend_per_share", units=("USD/shares",)))
    proceeds_annual = _annual_rows(rows("proceeds_debt"))
    repayments_annual = _annual_rows(rows("repayments_debt"))
    shares_rows = rows("shares_outstanding", units=("shares",))
    shares_annual = _annual_instant_rows(shares_rows)

    wb = openpyxl.Workbook()
    fin = wb.active
    fin.title = "Estados financieros"
    fin["A1"] = f"{submissions.get('name', ticker)} ({ticker}) -- SEC EDGAR (10-K), en millones de USD"
    fin["A1"].font = Font(bold=True, size=13)
    fin["A2"] = "Las filas en celeste se calculan por formula a partir de las filas de arriba -- corregi un numero y el ratio se recalcula solo."

    # --- Tabla A: ingresos, EBIT y margen (hasta 10 anios) ---
    row = 4
    fin.cell(row=row, column=2, value="Anio fiscal (cierre)").font = BOLD
    for i, end in enumerate(long_ends):
        c = fin.cell(row=row, column=3 + i, value=end)
        c.font = BOLD
        c.fill = FILL_HEADER
        c.font = HEADER_FONT
    row_revenue = row + 1
    _write_row(fin, row_revenue, "Ingresos", long_ends, revenue_by_end)
    row_ebit = row_revenue + 1
    _write_row(fin, row_ebit, "EBIT", long_ends, ebit_by_end)
    row_margin = row_ebit + 1
    fin.cell(row=row_margin, column=2, value="Margen EBIT")
    row_growth = row_margin + 1
    fin.cell(row=row_growth, column=2, value="Crecimiento de ingresos (a/a)")
    for i, end in enumerate(long_ends):
        col = get_column_letter(3 + i)
        mcell = fin.cell(row=row_margin, column=3 + i,
                          value=f"=IF({col}{row_revenue}=\"\",\"\",{col}{row_ebit}/{col}{row_revenue})")
        mcell.number_format = "0.0%"
        mcell.fill = FILL_FORMULA
        if i > 0:
            prev_col = get_column_letter(3 + i - 1)
            gcell = fin.cell(row=row_growth, column=3 + i,
                              value=f"=IF(OR({col}{row_revenue}=\"\",{prev_col}{row_revenue}=\"\"),\"\","
                                    f"{col}{row_revenue}/{prev_col}{row_revenue}-1)")
            gcell.number_format = "0.0%"
            gcell.fill = FILL_FORMULA

    margin_range = f"{get_column_letter(3)}{row_margin}:{get_column_letter(2 + len(long_ends))}{row_margin}"
    margin_range_3y = f"{get_column_letter(2 + len(long_ends) - 2)}{row_margin}:{get_column_letter(2 + len(long_ends))}{row_margin}"
    margin_range_5y = f"{get_column_letter(max(3, 2 + len(long_ends) - 4))}{row_margin}:{get_column_letter(2 + len(long_ends))}{row_margin}"

    # --- Tabla B: detalle anual (hasta 5 anios) ---
    row_b_header = row_growth + 3
    fin.cell(row=row_b_header, column=2, value="Anio fiscal (cierre)").font = BOLD
    for i, end in enumerate(short_ends):
        c = fin.cell(row=row_b_header, column=3 + i, value=end)
        c.fill = FILL_HEADER
        c.font = HEADER_FONT

    def _short_vals(annual_rows, unit_div=1_000_000):
        by_end = {r["end"]: r["val"] / unit_div for r in annual_rows}
        return {e: by_end.get(e) for e in short_ends}

    r = row_b_header + 1
    row_interest = r; _write_row(fin, r, "Gasto en intereses", short_ends, _short_vals(interest_annual)); r += 1
    row_da = r; _write_row(fin, r, "D&A", short_ends, _short_vals(da_annual)); r += 1
    row_capex = r; _write_row(fin, r, "CapEx", short_ends, _short_vals(capex_annual)); r += 1
    row_rd = r; _write_row(fin, r, "Gasto en I+D", short_ends, _short_vals(rd_annual)); r += 1
    row_tax = r; _write_row(fin, r, "Impuesto a las ganancias", short_ends, _short_vals(tax_annual)); r += 1
    row_pretax = r; _write_row(fin, r, "Resultado antes de impuestos", short_ends, _short_vals(pretax_annual)); r += 1
    row_div = r; _write_row(fin, r, "Dividendo por accion (USD)", short_ends, _short_vals(dividend_annual, 1),
                             number_format="#,##0.00"); r += 1
    row_proceeds = r; _write_row(fin, r, "Emision de deuda", short_ends, _short_vals(proceeds_annual)); r += 1
    row_repay = r; _write_row(fin, r, "Repago de deuda", short_ends, _short_vals(repayments_annual)); r += 1
    row_shares = r; _write_row(fin, r, "Acciones en circulacion (cierre, M)", short_ends,
                                _short_vals(shares_annual)); r += 1

    r += 1
    row_ratio_interest = r
    fin.cell(row=r, column=2, value="Interes / EBIT"); r += 1
    row_ratio_da = r
    fin.cell(row=r, column=2, value="D&A / Ingresos"); r += 1
    row_ratio_capex = r
    fin.cell(row=r, column=2, value="CapEx / Ingresos"); r += 1
    row_ratio_tax = r
    fin.cell(row=r, column=2, value="Tasa impositiva efectiva"); r += 1
    row_ratio_borrow = r
    fin.cell(row=r, column=2, value="Endeudamiento neto / Ingresos"); r += 1
    row_ratio_div_growth = r
    fin.cell(row=r, column=2, value="Crecimiento de dividendos (a/a)"); r += 1
    row_ratio_shares_growth = r
    fin.cell(row=r, column=2, value="Crecimiento de acciones (a/a)"); r += 1

    revenue_short_row_refs = [get_column_letter(3 + i) + str(row_revenue) for i, e in enumerate(long_ends) if e in short_ends]
    # mapear cada columna corta (Tabla B) a su columna correspondiente de ingresos/EBIT en la Tabla A
    long_col_by_end = {e: get_column_letter(3 + i) for i, e in enumerate(long_ends)}

    for i, end in enumerate(short_ends):
        col = get_column_letter(3 + i)
        rev_col = long_col_by_end[end]
        cell = fin.cell(row=row_ratio_interest, column=3 + i,
                         value=f"=IF(OR({col}{row_interest}=\"\",{rev_col}{row_ebit}=\"\"),\"\",{col}{row_interest}/{rev_col}{row_ebit})")
        cell.number_format = "0.0%"; cell.fill = FILL_FORMULA
        cell = fin.cell(row=row_ratio_da, column=3 + i,
                         value=f"=IF(OR({col}{row_da}=\"\",{rev_col}{row_revenue}=\"\"),\"\",{col}{row_da}/{rev_col}{row_revenue})")
        cell.number_format = "0.0%"; cell.fill = FILL_FORMULA
        cell = fin.cell(row=row_ratio_capex, column=3 + i,
                         value=f"=IF(OR({col}{row_capex}=\"\",{rev_col}{row_revenue}=\"\"),\"\",{col}{row_capex}/{rev_col}{row_revenue})")
        cell.number_format = "0.0%"; cell.fill = FILL_FORMULA
        cell = fin.cell(row=row_ratio_tax, column=3 + i,
                         value=f"=IF(OR({col}{row_tax}=\"\",{col}{row_pretax}=\"\"),\"\",{col}{row_tax}/{col}{row_pretax})")
        cell.number_format = "0.0%"; cell.fill = FILL_FORMULA
        cell = fin.cell(row=row_ratio_borrow, column=3 + i,
                         value=f"=IF(OR({col}{row_proceeds}=\"\",{col}{row_repay}=\"\",{rev_col}{row_revenue}=\"\"),\"\","
                               f"({col}{row_proceeds}-{col}{row_repay})/{rev_col}{row_revenue})")
        cell.number_format = "0.0%"; cell.fill = FILL_FORMULA
        if i > 0:
            prev_col = get_column_letter(3 + i - 1)
            cell = fin.cell(row=row_ratio_div_growth, column=3 + i,
                             value=f"=IF(OR({col}{row_div}=\"\",{prev_col}{row_div}=0,{prev_col}{row_div}=\"\"),\"\","
                                   f"{col}{row_div}/{prev_col}{row_div}-1)")
            cell.number_format = "0.0%"; cell.fill = FILL_FORMULA
            cell = fin.cell(row=row_ratio_shares_growth, column=3 + i,
                             value=f"=IF(OR({col}{row_shares}=\"\",{prev_col}{row_shares}=\"\"),\"\","
                                   f"{col}{row_shares}/{prev_col}{row_shares}-1)")
            cell.number_format = "0.0%"; cell.fill = FILL_FORMULA

    ratio_da_range = f"C{row_ratio_da}:{get_column_letter(2 + len(short_ends))}{row_ratio_da}"
    ratio_capex_range = f"C{row_ratio_capex}:{get_column_letter(2 + len(short_ends))}{row_ratio_capex}"
    ratio_interest_range = f"C{row_ratio_interest}:{get_column_letter(2 + len(short_ends))}{row_ratio_interest}"
    ratio_tax_range = f"C{row_ratio_tax}:{get_column_letter(2 + len(short_ends))}{row_ratio_tax}"
    ratio_borrow_range = f"C{row_ratio_borrow}:{get_column_letter(2 + len(short_ends))}{row_ratio_borrow}"

    # --- Tabla C: LTM (suma de los ultimos 4 trimestres, si son consecutivos) y balance ---
    row_c = row_ratio_shares_growth + 3
    fin.cell(row=row_c, column=2,
              value="LTM (suma de 4 trimestres consecutivos; si faltan o hay un hueco, se usa el ultimo 10-K)").font = BOLD
    q_revenue = _last_four_contiguous_quarters(rows("revenue"))
    q_ebit = _last_four_contiguous_quarters(rows("ebit"))
    q_interest = _last_four_contiguous_quarters(rows("interest_expense"))

    def _write_quarters(row_label, label, qrows):
        """Escribe hasta 4 trimestres en C..F y su suma en H (columna 8) --
        si qrows viene vacio (sin 4 trimestres consecutivos), no escribe nada
        y el llamador usa el ultimo 10-K como referencia en su lugar."""
        fin.cell(row=row_label, column=2, value=label)
        if not qrows:
            return
        for i, qr in enumerate(qrows):
            c = fin.cell(row=row_label, column=3 + i, value=qr["val"] / 1_000_000)
            c.number_format = "#,##0.0"
        last_col = get_column_letter(2 + len(qrows))
        total_cell = fin.cell(row=row_label, column=8, value=f"=SUM(C{row_label}:{last_col}{row_label})")
        total_cell.number_format = "#,##0.0"
        total_cell.fill = FILL_FORMULA

    r = row_c + 1
    fin.cell(row=r, column=2, value="(columnas C-F = ultimos 4 trimestres; columna H = suma = LTM)")
    r += 1
    row_ltm_revenue = r; _write_quarters(r, "Ingresos (trimestres)", q_revenue); r += 1
    row_ltm_ebit = r; _write_quarters(r, "EBIT (trimestres)", q_ebit); r += 1
    row_ltm_interest = r; _write_quarters(r, "Intereses (trimestres)", q_interest); r += 1

    r += 2
    fin.cell(row=r, column=2, value="Balance -- ultimo 10-K cerrado vs. mas reciente (10-Q/10-K), en millones").font = BOLD
    r += 1
    fin.cell(row=r, column=3, value="Ultimo 10-K").font = BOLD
    fin.cell(row=r, column=4, value="Mas reciente").font = BOLD
    r += 1

    equity_rows = rows("equity")
    lt_debt_rows, cur_debt_rows = rows("long_term_debt"), rows("current_debt")
    cash_rows = rows("cash")
    minority_rows = rows("minority_interest")
    nol_rows = rows("nol")

    def _balance_row(label, concept_rows, as_of_fn_needed=True):
        nonlocal r
        fin.cell(row=r, column=2, value=label)
        prior = _instant_as_of(concept_rows, last_10k_end) if concept_rows else None
        latest = _latest_instant(concept_rows) if concept_rows else None
        fin.cell(row=r, column=3, value=(prior or 0) / 1_000_000).number_format = "#,##0.0"
        fin.cell(row=r, column=4, value=(latest or 0) / 1_000_000).number_format = "#,##0.0"
        row_here = r
        r += 1
        return row_here

    row_bal_equity = _balance_row("Valor libro del equity", equity_rows)
    row_bal_debt_lt = _balance_row("Deuda de largo plazo", lt_debt_rows)
    row_bal_debt_cur = _balance_row("Deuda corriente", cur_debt_rows)
    r_debt_total = r
    fin.cell(row=r_debt_total, column=2, value="Deuda total (LP + corriente)")
    fin.cell(row=r_debt_total, column=3, value=f"=C{row_bal_debt_lt}+C{row_bal_debt_cur}").number_format = "#,##0.0"
    fin.cell(row=r_debt_total, column=4, value=f"=D{row_bal_debt_lt}+D{row_bal_debt_cur}").number_format = "#,##0.0"
    fin.cell(row=r_debt_total, column=3).fill = FILL_FORMULA
    fin.cell(row=r_debt_total, column=4).fill = FILL_FORMULA
    r += 1
    row_bal_cash = _balance_row("Caja", cash_rows)
    row_bal_minority = _balance_row("Interes minoritario", minority_rows)
    row_bal_nol = _balance_row("NOL acumulado", nol_rows)
    row_bal_shares = r
    fin.cell(row=r, column=2, value="Acciones en circulacion (M)")
    prior_shares = _instant_as_of(shares_rows, last_10k_end)
    latest_shares = _latest_instant(shares_rows)
    fin.cell(row=r, column=3, value=(prior_shares or 0) / 1_000_000).number_format = "#,##0.0"
    fin.cell(row=r, column=4, value=(latest_shares or 0) / 1_000_000).number_format = "#,##0.0"
    r += 1

    for col in range(1, 3 + len(short_ends) + 2):
        fin.column_dimensions[get_column_letter(col)].width = 16
    fin.column_dimensions["B"].width = 34

    # ================= Hoja 2: Inputs dashboard =================
    inp = wb.create_sheet("Inputs dashboard")
    inp["A1"] = "Campo (mismo nombre que en el dashboard)"
    inp["B1"] = "Valor"
    inp["A1"].font = BOLD
    inp["B1"].font = BOLD
    inp["A2"] = "Celeste = calculado por formula desde 'Estados financieros'. Amarillo = SEC EDGAR no lo tiene, completalo vos."
    inp["A2"].font = Font(italic=True, size=9)

    country = _resolve_country(submissions)
    industry = submissions.get("sicDescription") or ""

    def put(row_i, field, value, formula=False, manual=False, number_format=None):
        inp.cell(row=row_i, column=1, value=field)
        c = inp.cell(row=row_i, column=2, value=value)
        if formula:
            c.fill = FILL_FORMULA
        if manual:
            c.fill = FILL_INPUT
        if number_format:
            c.number_format = number_format
        return row_i + 1

    # Referencias SIN el prefijo de hoja -- se les antepone "'Estados
    # financieros'!" uniformemente donde se usan, asi que si aca ya llevaran
    # el prefijo (caso sin trimestres validos) quedaria duplicado.
    last_short_col = get_column_letter(2 + len(short_ends))
    last_long_col = get_column_letter(2 + len(long_ends))
    ltm_rev_ref = f"H{row_ltm_revenue}" if q_revenue else f"{last_long_col}{row_revenue}"
    ltm_ebit_ref = f"H{row_ltm_ebit}" if q_ebit else f"{last_long_col}{row_ebit}"
    ltm_int_ref = f"H{row_ltm_interest}" if q_interest else f"{last_short_col}{row_interest}"

    i = 3
    i = put(i, "ticker", ticker.upper())
    i = put(i, "company_name", submissions.get("name") or ticker)
    i = put(i, "country_of_incorporation", country)
    i = put(i, "industry_us", industry)
    i = put(i, "industry_global", industry)
    i = put(i, "current_price", current_price or None, manual=(not current_price))
    i = put(i, "revenue_ltm", f"='Estados financieros'!{ltm_rev_ref}", formula=True, number_format="#,##0.0")
    i = put(i, "revenue_prior_10k", f"='Estados financieros'!{last_long_col}{row_revenue}", formula=True, number_format="#,##0.0")
    i = put(i, "ebit_ltm", f"='Estados financieros'!{ltm_ebit_ref}", formula=True, number_format="#,##0.0")
    i = put(i, "ebit_prior_10k", f"='Estados financieros'!{last_long_col}{row_ebit}", formula=True, number_format="#,##0.0")
    i = put(i, "interest_expense_ltm", f"='Estados financieros'!{ltm_int_ref}", formula=True, number_format="#,##0.0")
    i = put(i, "interest_expense_prior_10k", f"='Estados financieros'!{last_short_col}{row_interest}", formula=True, number_format="#,##0.0")
    i = put(i, "book_value_equity_ltm", f"='Estados financieros'!D{row_bal_equity}", formula=True, number_format="#,##0.0")
    i = put(i, "book_value_equity_prior_10k", f"='Estados financieros'!C{row_bal_equity}", formula=True, number_format="#,##0.0")
    i = put(i, "book_value_debt_ltm", f"='Estados financieros'!D{r_debt_total}", formula=True, number_format="#,##0.0")
    i = put(i, "book_value_debt_prior_10k", f"='Estados financieros'!C{r_debt_total}", formula=True, number_format="#,##0.0")
    i = put(i, "cash_ltm", f"='Estados financieros'!D{row_bal_cash}", formula=True, number_format="#,##0.0")
    i = put(i, "cash_prior_10k", f"='Estados financieros'!C{row_bal_cash}", formula=True, number_format="#,##0.0")
    i = put(i, "minority_interests", f"='Estados financieros'!D{row_bal_minority}", formula=True, number_format="#,##0.0")
    i = put(i, "shares_outstanding", f"='Estados financieros'!D{row_bal_shares}", formula=True, number_format="#,##0.0")
    i = put(i, "nol_carryforward", f"='Estados financieros'!D{row_bal_nol}", formula=True, number_format="#,##0.0")
    i = put(i, "effective_tax_rate", f"=AVERAGE('Estados financieros'!{ratio_tax_range})",
            formula=True, number_format="0.0%")
    i = put(i, "dividend_per_share_ltm", f"='Estados financieros'!{last_short_col}{row_div}", formula=True, number_format="#,##0.00")
    i = put(i, "ebit_margin_ltm",
            f"='Estados financieros'!{ltm_ebit_ref}/'Estados financieros'!{ltm_rev_ref}",
            formula=True, number_format="0.0%")
    i = put(i, "ebit_margin_avg_3y", f"=AVERAGE('Estados financieros'!{margin_range_3y})", formula=True, number_format="0.0%")
    i = put(i, "ebit_margin_avg_5y", f"=AVERAGE('Estados financieros'!{margin_range_5y})", formula=True, number_format="0.0%")
    i = put(i, "ebit_margin_avg_10y", f"=AVERAGE('Estados financieros'!{margin_range})", formula=True, number_format="0.0%")
    revenue_growth_formula = _cagr_formula("Estados financieros", row_revenue, last_long_col, 3, "C")
    i = put(i, "revenue_growth_next_year", revenue_growth_formula, formula=True, number_format="0.0%")
    i = put(i, "revenue_growth_years_2_to_5", revenue_growth_formula, formula=True, number_format="0.0%")
    i = put(i, "riskfree_rate", riskfree_rate or None, manual=(not riskfree_rate), number_format="0.00%")
    i = put(i, "initial_cost_of_capital", initial_cost_of_capital or None, manual=(not initial_cost_of_capital), number_format="0.00%")
    i = put(i, "hist_interest_pct_of_ebit", f"=AVERAGE('Estados financieros'!{ratio_interest_range})", formula=True, number_format="0.0%")
    i = put(i, "hist_da_pct_of_revenue", f"=AVERAGE('Estados financieros'!{ratio_da_range})", formula=True, number_format="0.0%")
    i = put(i, "hist_capex_pct_of_revenue", f"=AVERAGE('Estados financieros'!{ratio_capex_range})", formula=True, number_format="0.0%")
    i = put(i, "hist_net_borrowing_pct_of_revenue", f"=AVERAGE('Estados financieros'!{ratio_borrow_range})", formula=True, number_format="0.0%")
    i = put(i, "hist_shares_growth_rate", _cagr_formula("Estados financieros", row_shares, last_short_col, 3, "C"),
            formula=True, number_format="0.0%")
    i = put(i, "hist_dividend_growth_rate", _cagr_formula("Estados financieros", row_div, last_short_col, 3, "C"),
            formula=True, number_format="0.0%")
    i = put(i, "capitalize_rd", "Yes" if (rd_annual and rd_annual[-1]["val"]) else "No")
    i = put(i, "rd_expense_current_year", f"='Estados financieros'!{last_short_col}{row_rd}", formula=True, number_format="#,##0.0")
    for k in range(1, 10):
        col_idx = len(short_ends) - 1 - k
        if col_idx >= 0:
            col = get_column_letter(3 + col_idx)
            i = put(i, f"rd_expense_year_minus_{k}", f"='Estados financieros'!{col}{row_rd}", formula=True, number_format="#,##0.0")
        else:
            i = put(i, f"rd_expense_year_minus_{k}", 0.0)

    inp.column_dimensions["A"].width = 34
    inp.column_dimensions["B"].width = 40

    return wb, submissions


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("ticker")
    parser.add_argument("--price", type=float, default=0.0)
    parser.add_argument("--riskfree", type=float, default=0.0)
    parser.add_argument("--wacc", type=float, default=0.0)
    parser.add_argument("--out-dir", default=str(Path(__file__).resolve().parent.parent / "data"))
    args = parser.parse_args(argv[1:])

    try:
        wb, submissions = build_workbook(args.ticker, args.price, args.riskfree, args.wacc)
    except SecEdgarError as exc:
        print(f"Error: {exc}")
        return 1

    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    xlsx_path = out_dir / f"{args.ticker.upper()}_sec_edgar.xlsx"
    wb.save(xlsx_path)
    print(f"Generado: {xlsx_path}")

    csv_path = out_dir / f"{args.ticker.upper()}_sec_edgar.csv"
    inputs = load_company_inputs_from_sec_edgar(
        args.ticker, current_price=args.price, riskfree_rate=args.riskfree,
        initial_cost_of_capital=args.wacc,
    )
    save_company_inputs(inputs, csv_path)
    print(f"Generado: {csv_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
