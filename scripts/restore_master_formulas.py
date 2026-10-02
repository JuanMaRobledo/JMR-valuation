#!/usr/bin/env python
"""Devuelve a una hoja del Modelo JMR las fórmulas de la plantilla maestra en las celdas que se cambiaron
después de una revisión de referencia (2-oct-2026, NKE).

Regla del usuario: todas las empresas usan las mismas fórmulas de la maestra; entre historias solo cambian los
supuestos (crecimiento inicial y objetivo, margen inicial y objetivo) y, por empresa, las celdas de entrada y los
datos reportados. El script compara la revisión de referencia (exportada a xlsx) con la versión actual, y para cada
celda distinta escribe el contenido de la referencia, salvo las celdas que se conservan (--keep). Antes de escribir
cada fórmula la contrasta con la de la plantilla maestra en la misma celda, cuando la maestra la tiene.

Uso:
    PYTHONPATH=.:scripts python scripts/restore_master_formulas.py --sheet-id ID --ref ref.xlsx --cur cur.xlsx \
        --keep "Input sheet!B24" --keep-sheet "Cash Flow Statement" --keep-sheet "Descuento de múltiplos" [--apply]
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
from pathlib import Path

import openpyxl

from jmr_valuation.io.sheets_auth import get_gspread_client

MASTER_ID = "19PRUFiYsNavUcN6WwHBVlp-VRMozp3rNSE2R1zt7N-g"
_ROOT = Path(__file__).resolve().parent.parent


def to_locale(v):
    """Fórmula de xlsx (separador ',' y decimal '.') a la de la hoja (es: ';' y ',')."""
    if not (isinstance(v, str) and v.startswith("=")):
        return v
    parts = v.split('"')
    for i in range(0, len(parts), 2):
        p = re.sub(r"(?<=\d)\.(?=\d)", "#DEC#", parts[i])
        parts[i] = p.replace(",", ";").replace("#DEC#", ",")
    return '"'.join(parts)


def norm(v):
    return re.sub(r"\s+", "", str(v)) if v is not None else ""


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--sheet-id", required=True)
    ap.add_argument("--ref", required=True)
    ap.add_argument("--cur", required=True)
    ap.add_argument("--keep", action="append", default=[])
    ap.add_argument("--keep-sheet", action="append", default=[])
    ap.add_argument("--apply", action="store_true")
    a = ap.parse_args()
    ref, cur = openpyxl.load_workbook(a.ref), openpyxl.load_workbook(a.cur)
    changes: dict[str, dict[str, object]] = {}
    for n in cur.sheetnames:
        if n not in ref.sheetnames or n in a.keep_sheet:
            continue
        wr, wc = ref[n], cur[n]
        rows, cols = max(wr.max_row, wc.max_row), max(wr.max_column, wc.max_column)
        for r in range(1, rows + 1):
            for c in range(1, cols + 1):
                addr = wc.cell(r, c).coordinate
                if f"{n}!{addr}" in a.keep:
                    continue
                vr, vc = wr.cell(r, c).value, wc.cell(r, c).value
                if vr != vc:
                    changes.setdefault(n, {})[addr] = vr
    client = get_gspread_client()
    sh, master = client.open_by_key(a.sheet_id), client.open_by_key(MASTER_ID)
    mtitles = {w.title for w in master.worksheets()}
    data, report, backup = [], {}, {}
    for n, cells in changes.items():
        q = "'" + n.replace("'", "''") + "'"
        curvals = sh.values_batch_get([f"{q}!{c}" for c in cells], params={"valueRenderOption": "FORMULA"})["valueRanges"]
        mvals = (master.values_batch_get([f"{q}!{c}" for c in cells], params={"valueRenderOption": "FORMULA"})["valueRanges"]
                 if n in mtitles else [{}] * len(cells))
        same_master = diff_master = 0
        for (addr, v), cv, mv in zip(cells.items(), curvals, mvals):
            new = to_locale(v) if v is not None else ""
            old = (cv.get("values") or [[""]])[0][0]
            backup.setdefault(n, {})[addr] = old
            mf = (mv.get("values") or [[None]])[0][0] if mv else None
            if isinstance(new, str) and new.startswith("=") and isinstance(mf, str) and mf.startswith("="):
                if norm(mf) == norm(new):
                    same_master += 1
                else:
                    diff_master += 1
            data.append({"range": f"{q}!{addr}", "values": [[new]]})
        report[n] = {"celdas": len(cells), "formula_igual_maestra": same_master, "formula_distinta_maestra": diff_master}
    print(json.dumps(report, ensure_ascii=False, indent=1))
    if a.apply:
        bp = _ROOT / "reference" / "backups" / f"restaurar_maestra_{a.sheet_id}_{dt.date.today().isoformat()}.json"
        bp.write_text(json.dumps(backup, ensure_ascii=False, indent=1, default=str))
        for i in range(0, len(data), 400):
            sh.values_batch_update({"valueInputOption": "USER_ENTERED", "data": data[i:i + 400]})
        print("escritas", len(data), "respaldo", bp)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
