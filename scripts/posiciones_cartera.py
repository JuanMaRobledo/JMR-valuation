#!/usr/bin/env python
"""Posiciones de la cartera para el bloque «Mi posición» de las hojas (6-oct-2026).

Escribe reference/cartera_compras_<hoy>.json con, por empresa valorada: si está en cartera, acciones, costo promedio,
primera compra de la posición vigente (la primera compra después de la última vez que la cantidad llegó a cero) y su
precio, las fechas de compra y las cuentas. Lo usa scripts/posicion_cartera.py.

Fuentes, en este orden:
1. La app Cartera (fuente principal): GET <CARTERA_URL>/api/export con "Authorization: Bearer $CARTERA_EXPORT_TOKEN".
   Se usa sola si las dos variables de entorno existen; --app la exige.
2. --json <archivo>: el JSON «Cartera completa» descargado del menú «Exportar cartera» de la app.
3. --xlsx <archivo>: la hoja «Seguimiento de cartera» (Drive 1P3yzgr-RXJU6fFGuVPw0_JLQMwh_khGWuAZsZmft7kM) exportada a
   .xlsx con el conector de Google Drive, o lectura directa con la cuenta de servicio si la hoja está compartida como
   lector con jmr-valuation-bot@jmr-valuation-508617.iam.gserviceaccount.com.

Uso: PYTHONPATH=.:scripts python scripts/posiciones_cartera.py [--app | --json archivo.json | --xlsx archivo.xlsx]
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import sys
import urllib.request
from collections import defaultdict
from pathlib import Path

_ROOT = Path(__file__).resolve().parents[1]
for p in (str(_ROOT), str(_ROOT / "scripts")):
    if p not in sys.path:
        sys.path.insert(0, p)

HOJA = "1P3yzgr-RXJU6fFGuVPw0_JLQMwh_khGWuAZsZmft7kM"
ALIAS = {"GOOG": {"GOOG", "GOOGL"}}


def filas(args) -> tuple[list, list, str]:
    if args.xlsx:
        import openpyxl
        wb = openpyxl.load_workbook(args.xlsx, data_only=True)
        return (list(wb["Transacciones"].iter_rows(values_only=True)), list(wb["Posiciones"].iter_rows(values_only=True)),
                f"«Seguimiento de cartera» (Drive, descargada {dt.date.today()})")
    import model_steps as ms
    sh = ms.open_sheet(HOJA)
    conv = lambda v: v  # noqa: E731
    tx = sh.worksheet("Transacciones").get_all_values(value_render_option="UNFORMATTED_VALUE")
    po = sh.worksheet("Posiciones").get_all_values(value_render_option="UNFORMATTED_VALUE")
    return [[conv(x) for x in r] for r in tx], [[conv(x) for x in r] for r in po], "«Seguimiento de cartera» (Drive, lectura directa)"


def fecha(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)):
        return dt.date(1899, 12, 30) + dt.timedelta(days=int(v))
    return dt.date.fromisoformat(str(v)[:10])


def export_app() -> dict:
    url = os.environ.get("CARTERA_URL", "").rstrip("/")
    token = os.environ.get("CARTERA_EXPORT_TOKEN", "")
    if not url or not token:
        raise SystemExit("Faltan CARTERA_URL y CARTERA_EXPORT_TOKEN en el entorno")
    req = urllib.request.Request(f"{url}/api/export", headers={"Authorization": f"Bearer {token}"})
    with urllib.request.urlopen(req, timeout=60) as r:
        return json.load(r)


def desde_app(data: dict, empresas: list[str]) -> dict:
    """Posiciones por empresa a partir de la exportación de la app (posiciones abiertas consolidadas)."""
    cuentas = {c["id"]: c["name"] for c in data["accounts"]}
    abiertas = defaultdict(list)
    for p in data["positions"]:
        abiertas[p["ticker"].upper()].append(p)
    cerradas = {p["ticker"].upper(): p for p in data.get("closedPositions", [])}
    hoy = data["generatedAt"][:10]
    out = {}
    for t in empresas:
        syms = ALIAS.get(t, {t})
        ps = [p for s in syms for p in abiertas.get(s, []) if p["quantity"] > 1e-9]
        if not ps:
            c = [cerradas[s] for s in syms if s in cerradas]
            out[t] = {"en_cartera": False,
                      "nota": f"sin posición al {hoy}" + (" (posición cerrada)" if c else " (sin operaciones)")}
            continue
        q = sum(p["quantity"] for p in ps)
        costo = sum(p["costBasisLocal"] for p in ps)
        ini = min(ps, key=lambda p: p["openSince"] or "9999")
        out[t] = {"en_cartera": True, "primera_compra": ini["openSince"],
                  "precio_primera_compra": round(ini["openSincePrice"], 4) if ini["openSincePrice"] else None,
                  "cantidad": round(q, 5), "costo_promedio": round(costo / q, 4),
                  "compras": sorted({d for p in ps for d in p["buyDates"]}),
                  "cuentas": sorted({cuentas.get(i, i) for p in ps for i in p["accountIds"]})}
    return out


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--app", action="store_true", help="exige leer de la app Cartera")
    ap.add_argument("--json", help="JSON «Cartera completa» descargado de la app")
    ap.add_argument("--xlsx")
    a = ap.parse_args(argv)
    empresas = sorted(json.loads((_ROOT / "reference" / "cartera_drive.json").read_text())["empresas"])
    hay_app = bool(os.environ.get("CARTERA_URL") and os.environ.get("CARTERA_EXPORT_TOKEN"))
    if a.app or a.json or (hay_app and not a.xlsx):
        data = json.loads(Path(a.json).read_text()) if a.json else export_app()
        fuente = (f"app Cartera ({'archivo exportado' if a.json else 'GET /api/export'}, "
                  f"generado {data['generatedAt'][:16].replace('T', ' ')} UTC)")
        return escribir(desde_app(data, empresas), fuente)
    tx, po, fuente = filas(a)
    hdr = list(tx[0])
    i = {k: hdr.index(k) for k in ("Fecha", "Cuenta", "Símbolo", "Actividad", "Cantidad", "Precio USD")}
    mov = defaultdict(list)
    for r in tx[1:]:
        if r and r[i["Actividad"]] in ("Compra", "Venta") and r[i["Símbolo"]]:
            mov[r[i["Símbolo"]]].append((fecha(r[i["Fecha"]]), r[i["Cuenta"]], r[i["Actividad"]], float(r[i["Cantidad"]]),
                                         float(r[i["Precio USD"]] or 0)))
    ph = next(k for k, r in enumerate(po) if r and r[0] == "Cuenta")
    hp = list(po[ph])
    pos = defaultdict(lambda: [0.0, 0.0])
    for r in po[ph + 1:]:
        if r and r[1]:
            pos[r[1]][0] += float(r[hp.index("Cantidad")] or 0)
            pos[r[1]][1] += float(r[hp.index("Costo base USD")] or 0)
    out = {}
    for t in empresas:
        syms = ALIAS.get(t, {t})
        L = sorted((x for s in syms for x in mov.get(s, [])), key=lambda x: x[0])
        q = sum(pos[s][0] for s in syms if s in pos)
        c = sum(pos[s][1] for s in syms if s in pos)
        if q <= 1e-9 or not L:
            ult = f" (última operación {L[-1][0]} {L[-1][2].lower()})" if L else " (sin operaciones)"
            out[t] = {"en_cartera": False, "nota": f"sin posición al {dt.date.today()}{ult}"}
            continue
        run, ini = 0.0, None
        for d, acc, act, qq, p in L:
            if run <= 1e-9 and act == "Compra":
                ini = d
            run += qq
        dia = [(qq, p) for d, acc, act, qq, p in L if d == ini and act == "Compra"]
        out[t] = {"en_cartera": True, "primera_compra": str(ini),
                  "precio_primera_compra": round(sum(q_ * p for q_, p in dia) / sum(q_ for q_, p in dia), 4),
                  "cantidad": round(q, 5), "costo_promedio": round(c / q, 4),
                  "compras": sorted({str(d) for d, acc, act, qq, p in L if act == "Compra" and d >= ini}),
                  "cuentas": sorted({acc for d, acc, act, qq, p in L if d >= ini})}
    return escribir(out, fuente)


def escribir(out: dict, fuente: str) -> int:
    dest = _ROOT / "reference" / f"cartera_compras_{dt.date.today()}.json"
    dest.write_text(json.dumps({"fecha": str(dt.date.today()), "fuente": fuente, "empresas": out}, ensure_ascii=False,
                               indent=1) + "\n")
    for t, v in out.items():
        print(t, f"{v['cantidad']:g} acc. a {v['costo_promedio']}" if v["en_cartera"] else v["nota"])
    print("escrito", dest)
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
