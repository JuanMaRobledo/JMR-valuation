import openpyxl
import pytest

from jmr_valuation.io.excel_loader import load_company_inputs_from_excel


def _build_minimal_workbook(tmp_path, j19_overrides: dict[str, float] | None = None):
    """Arma un .xlsx sintetico minimo con las hojas/celdas que lee excel_loader,
    simulando que Excel ya lo recalculo (los valores estan "cacheados" porque
    openpyxl los escribe como valores literales, no formulas)."""
    wb = openpyxl.Workbook()
    ins = wb.active
    ins.title = "Input sheet"
    values = {
        "B1": "TEST", "B5": "Test Co", "B8": "United States", "B9": "Software (System & Application)",
        "B10": "Software (System & Application)",
        "B12": 1000, "C12": 900, "D12": 1.0, "B13": 200, "C13": 180,
        "B14": 5, "C14": 4, "B15": 500, "C15": 480, "B16": 300, "C16": 280,
        "B17": "No", "B18": "No", "B19": 100, "C19": 90, "B20": 10, "C20": 8,
        "B21": 0, "B22": 50, "B23": 40, "B24": 0.2, "B25": 0.25,
        "B27": 0.10, "B29": 0.10, "B30": 0.25, "B31": 5, "B32": 2.0, "B33": 2.0,
        "B35": 0.04, "B36": 0.09, "B38": "No", "B39": 0, "B40": 0, "B41": 0, "B42": 0,
        "B63": 0,
    }
    for coord, v in values.items():
        ins[coord] = v

    resumen = wb.create_sheet("Resumen de Valoración")
    resumen["G3"] = "Software"
    resumen["G4"] = 0.30

    tv = wb.create_sheet("Trailing Valuation")
    for row in (13, 15, 16, 21, 24):
        tv.cell(row=row, column=9, value=20 + row * 0.1)   # I (anio -3)
        tv.cell(row=row, column=10, value=22 + row * 0.1)  # J (anio -2)
        tv.cell(row=row, column=11, value=24 + row * 0.1)  # K (anio -1)

    inc = wb.create_sheet("Income Statement")
    for row, base in [(3, 800), (12, 150), (18, 5), (26, 45)]:
        for col, val in zip((8, 9, 10, 11), (base, base * 1.05, base * 1.10, base * 1.15)):
            inc.cell(row=row, column=col, value=val)

    cf = wb.create_sheet("Cash Flow Statement")
    for row, base in [(4, 40), (15, -30), (25, 0), (28, 5)]:
        for col, val in zip((8, 9, 10, 11), (base, base * 1.05, base * 1.10, base * 1.15)):
            cf.cell(row=row, column=col, value=val)

    bs = wb.create_sheet("Balance Sheet")
    for row, base in [(5, 200), (10, 300), (20, 10), (21, 5), (24, 250)]:
        for col, val in zip((8, 9, 10, 11), (base, base * 1.02, base * 1.04, base * 1.06)):
            bs.cell(row=row, column=col, value=val)

    if j19_overrides:
        for sheet_name, value in j19_overrides.items():
            ws = wb.create_sheet(sheet_name)
            ws["J19"] = value

    path = tmp_path / "synthetic_model.xlsx"
    wb.save(path)
    return path


def test_load_company_inputs_from_excel_reads_input_sheet(tmp_path):
    path = _build_minimal_workbook(tmp_path)
    inputs = load_company_inputs_from_excel(path)

    assert inputs.ticker == "TEST"
    assert inputs.company_name == "Test Co"
    assert inputs.revenue_ltm == 1000
    assert inputs.ebit_ltm == 200
    assert inputs.capitalize_rd is False
    assert inputs.has_operating_leases is False
    assert inputs.shares_outstanding == 50


def test_load_company_inputs_from_excel_reads_blend_settings(tmp_path):
    path = _build_minimal_workbook(tmp_path)
    inputs = load_company_inputs_from_excel(path)
    assert inputs.company_type == "Software"
    assert inputs.margin_of_safety == pytest.approx(0.30)


def test_load_company_inputs_from_excel_reads_trailing_multiples(tmp_path):
    path = _build_minimal_workbook(tmp_path)
    inputs = load_company_inputs_from_excel(path)
    # fila 13 = P/E: I=21.3, J=23.3, K=25.3 (base + row*0.1 = 13*0.1=1.3)
    assert inputs.hist_multiple_pe_y3 == pytest.approx(21.3)
    assert inputs.hist_multiple_pe_y2 == pytest.approx(23.3)
    assert inputs.hist_multiple_pe_y1 == pytest.approx(25.3)


def test_load_company_inputs_from_excel_no_j19_means_no_override(tmp_path):
    path = _build_minimal_workbook(tmp_path)
    inputs = load_company_inputs_from_excel(path)
    assert inputs.override_ev_fcff is False
    assert inputs.override_pe is False


def test_load_company_inputs_from_excel_j19_becomes_the_override(tmp_path):
    path = _build_minimal_workbook(tmp_path, j19_overrides={"EVFCFF": 14.6, "PE": 18.5})
    inputs = load_company_inputs_from_excel(path)

    assert inputs.override_ev_fcff is True
    assert inputs.multiple_ev_fcff == pytest.approx(14.6)
    assert inputs.override_pe is True
    assert inputs.multiple_pe == pytest.approx(18.5)
    # Las que no tienen J19 en el Excel no deben quedar overriden.
    assert inputs.override_p_ocf is False


def test_load_company_inputs_from_excel_computes_historical_ratios(tmp_path):
    path = _build_minimal_workbook(tmp_path)
    inputs = load_company_inputs_from_excel(path)
    assert inputs.hist_da_pct_of_revenue > 0
    assert inputs.hist_capex_pct_of_revenue > 0


def test_load_company_inputs_from_excel_produces_a_valid_valuation(tmp_path):
    from jmr_valuation.valuation import run_valuation
    path = _build_minimal_workbook(tmp_path)
    inputs = load_company_inputs_from_excel(path)
    report = run_valuation(inputs)
    assert report.scenarios["Base"].equity_bridge.value_per_share > 0
